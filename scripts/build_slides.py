"""Build the predefense talk from slides_mds/{objective,methodology,slides_styles}.md.

Run from anywhere:
    conda run -n talkbuild python scripts/build_slides.py

Outputs slides/predefense.pptx.  Intermediate media (equation and reference
PNGs, the OpenFOAM logo PNG, the mesh-motion MP4) go to slides/media/.
The POD video comes from scripts/make_pod_video.py (flowkit env).
Needs pdflatex + pdftocairo (system) and ffmpeg + rsvg-convert (talkbuild).
"""
import copy
import subprocess
from pathlib import Path

import numpy as np
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from speaker_notes import pptx_text

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'slides'
MEDIA = OUT / 'media'
FRAMES = ROOT.parents[1] / 'mesh_motion_anime'   # ~/Research/Thesis/mesh_motion_anime

INK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x8B, 0x1E, 0x2D)
GREY = RGBColor(0x4A, 0x4F, 0x5A)
RULE = RGBColor(0xD0, 0xD4, 0xDC)
FONT = 'Calibri'

EQ_DPI = 600
EQ_SCALE = 2.0          # LaTeX 10 pt -> about 20 pt on the slide


# --------------------------------------------------------------- media
def run(cmd, **kw):
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **kw)


def latex_png(name, body, math=True):
    """Render a display equation (or, with math=False, a text block) with
    pdflatex and return the PNG path."""
    tex = MEDIA / f'{name}.tex'
    if math:
        body = '$\\displaystyle ' + body + '$'
    tex.write_text(
        '\\documentclass[border=2pt,preview]{standalone}\n'
        '\\usepackage[T1]{fontenc}\\usepackage{lmodern}\n'
        '\\usepackage{amsmath,amssymb,bm,xcolor}\n'
        '\\definecolor{accent}{RGB}{139,30,45}\n'
        '\\definecolor{ink}{RGB}{31,42,68}\n'
        '\\begin{document}\n' + body + '\n\\end{document}\n')
    run(['pdflatex', '-interaction=nonstopmode', '-halt-on-error', tex.name], cwd=MEDIA)
    run(['pdftocairo', '-r', str(EQ_DPI), '-png', '-transp', '-singlefile', f'{name}.pdf', name], cwd=MEDIA)
    for ext in ('aux', 'log', 'tex', 'pdf'):
        (MEDIA / f'{name}.{ext}').unlink(missing_ok=True)
    return MEDIA / f'{name}.png'


def openfoam_logo():
    png = MEDIA / 'OpenFOAM_logo.png'
    run(['rsvg-convert', '-w', '2000', '-b', 'white', '-o', str(png),
         str(ROOT / 'slides_mds' / 'OpenFOAM_logo.svg')])
    return png


def longest_run(ok):
    """[start, stop) of the longest run of True: the figure body, without borders or slivers."""
    best, start = (0, 0), None
    for i, v in enumerate(np.append(ok, False)):
        if v and start is None:
            start = i
        elif not v and start is not None:
            best = max(best, (start, i), key=lambda b: b[1] - b[0])
            start = None
    return best


def kurtulus_panels():
    """Figures from Kurtulus (2016): trim the white page border only, no other edits."""
    out = []
    for k in range(1, 6):
        im = Image.open(ROOT / 'assets' / 'kurtulus_modes' / f'mode{k}.png').convert('RGB')
        white = np.asarray(im).min(axis=2) > 225
        r0, r1 = longest_run(white.mean(axis=1) < 0.05)
        c0, c1 = longest_run(white.mean(axis=0) < 0.05)
        png = MEDIA / f'kurtulus_mode{k}.png'
        im.crop((c0 + 2, r0 + 2, c1 - 2, r1 - 2)).save(png)
        out.append(png)
    return out


DURANTE_PDF = ROOT / 'reference_sources' / 'Durante_CNSNS_2020_Airfoil-Bifurcations-Chaos.pdf'
# (page, x0, y0, x1, y1) in pixels at 100 dpi, one panel each from Figs 4-6, 8-10, 12-14
DURANTE_PANELS = {
    'p1': {'history': (11, 245, 168, 602, 287), 'spectrum': (12, 430, 172, 622, 300),
           'portrait': (12, 440, 518, 622, 650)},
    'p2': {'history': (11, 245, 440, 602, 560), 'spectrum': (12, 430, 325, 622, 448),
           'portrait': (12, 440, 668, 622, 800)},
    'p3': {'history': (14, 430, 222, 672, 302), 'spectrum': (14, 430, 558, 622, 685),
           'portrait': (15, 430, 172, 622, 300)},
    'chaos': {'history': (16, 180, 395, 668, 555), 'spectrum': (17, 425, 172, 695, 347),
              'portrait': (17, 425, 565, 695, 748)},
}


def durante_panels(dpi=400):
    """Crop single panels out of the paper's figures (rendered from the PDF, no other edits)."""
    k = dpi / 100
    out = {}
    for regime, rows in DURANTE_PANELS.items():
        for row, (page, x0, y0, x1, y1) in rows.items():
            stem = MEDIA / f'durante_{regime}_{row}'
            run(['pdftoppm', '-r', str(dpi), '-f', str(page), '-l', str(page), '-singlefile', '-png',
                 '-x', str(round(x0 * k)), '-y', str(round(y0 * k)),
                 '-W', str(round((x1 - x0) * k)), '-H', str(round((y1 - y0) * k)),
                 str(DURANTE_PDF), str(stem)])
            out[regime, row] = stem.with_suffix('.png')
    return out


def gupta_map(dpi=500):
    """Fig. 7(a) of Gupta et al. (2023), page 14, cropped from the PDF (box in px at 100 dpi)."""
    x0, y0, x1, y1 = 112, 82, 582, 322
    k = dpi / 100
    stem = MEDIA / 'gupta_regime_map'
    run(['pdftoppm', '-r', str(dpi), '-f', '14', '-l', '14', '-singlefile', '-png',
         '-x', str(round(x0 * k)), '-y', str(round(y0 * k)),
         '-W', str(round((x1 - x0) * k)), '-H', str(round((y1 - y0) * k)),
         str(ROOT / 'reference_sources' / 'gupta_2D3DWakeTransitionsNACA0012.pdf'), str(stem)])
    return stem.with_suffix('.png')


def mesh_motion_video():
    """75 ParaView frames -> H.264 MP4 at 15 fps, final pose held 1.5 s."""
    mp4 = MEDIA / 'mesh_motion.mp4'
    run(['ffmpeg', '-y', '-framerate', '15', '-i', str(FRAMES / 'moveMesh.%04d.png'),
         '-vf', 'tpad=stop_mode=clone:stop_duration=1.5',
         '-c:v', 'libx264', '-preset', 'slow', '-crf', '20', '-pix_fmt', 'yuv420p',
         '-movflags', '+faststart', str(mp4)])
    return mp4


# --------------------------------------------------------------- helpers
def textbox(slide, x, y, w, h, text, size=18, bold=False, color=INK,
            align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color
    return tb


SECTIONS = {}            # slide -> footer label; page numbers are added after all slides exist


def new_slide(prs, section):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    SECTIONS[s.slide_id] = section
    return s


def rule(slide, x, y, w=1.4):
    bar = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def header(slide, title, size=32, rule_y=1.1):
    textbox(slide, 0.6, 0.35, 12.1, rule_y - 0.3, title, size=size, bold=True)
    rule(slide, 0.6, rule_y)


def footers(prs):
    for n, s in enumerate(prs.slides, start=1):
        section = SECTIONS.get(s.slide_id)
        if section is None:                     # cover
            continue
        if section:
            textbox(s, 0.6, 7.0, 6, 0.35, section, size=12, color=GREY)
        textbox(s, 11.73, 7.0, 1.0, 0.35, str(n), size=14, bold=True, align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, items, size=20, gap=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.3)))
        pPr.set('indent', str(-Inches(0.3)))
        bu = etree.SubElement(pPr, qn('a:buClr'))
        etree.SubElement(bu, qn('a:srgbClr')).set('val', str(ACCENT))
        etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        r = p.add_run()
        r.text = item
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.color.rgb = INK
    return tb


def video(slide, path, poster, x, y, w, h):
    movie = slide.shapes.add_movie(str(path), Inches(x), Inches(y), Inches(w), Inches(h),
                                   poster_frame_image=str(poster), mime_type='video/mp4')
    click_to_play(slide, movie)
    return movie


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def equation(slide, path, x, y, scale=EQ_SCALE):
    """Place an equation PNG at its natural size times scale."""
    wpx, hpx = Image.open(path).size
    w, h = wpx / EQ_DPI * scale, hpx / EQ_DPI * scale
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return h


def table(slide, rows, x, y, w, col0, size=16, row_h=0.46):
    shp = slide.shapes.add_table(len(rows), 2, Inches(x), Inches(y),
                                 Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    # plain look: no banding or theme style
    tblPr = shp._element.graphic.graphicData.tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    tbl.columns[0].width = Inches(col0)
    tbl.columns[1].width = Inches(w - col0)
    for i, (k, v) in enumerate(rows):
        for j, txt in enumerate((k, v)):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF7) if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.12)
            tf = c.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ''
            r = tf.paragraphs[0].add_run()
            r.text = txt
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold = j == 0
            r.font.color.rgb = INK
    return shp


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


CLICK_TO_PLAY = '''
<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
 <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:seq concurrent="1" nextAc="seek">
   <p:cTn id="2" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">
    <p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cond></p:stCondLst>
    <p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync>
    <p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>
     <p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>
      <p:childTnLst><p:par><p:cTn id="5" presetID="2" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="clickEffect">
       <p:stCondLst><p:cond delay="0"/></p:stCondLst>
       <p:childTnLst><p:cmd type="call" cmd="togglePause"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst>
      </p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst>
   </p:cTn>
   <p:nextCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cond></p:nextCondLst>
  </p:seq>
  <p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
 </p:childTnLst></p:cTn></p:par></p:tnLst>
</p:timing>'''


def click_to_play(slide, movie):
    """Play/pause when the video itself is clicked; plays once, no loop, no autostart."""
    sld = slide._element
    old = sld.find(qn('p:timing'))
    if old is not None:
        sld.remove(old)
    sld.append(etree.fromstring(CLICK_TO_PLAY.format(spid=movie.shape_id)))


REFS = {
    'eldredge2019': r"J.\,D. Eldredge and A.\,R. Jones,\\"
                    r"``Leading-edge vortices: Mechanics and modeling,''\\"
                    r"\emph{Annual Review of Fluid Mechanics}, vol.~51, pp.~75--104, 2019.",
    'williamson': r"C.\,H.\,K. Williamson, ``Vortex dynamics in the cylinder wake,'' \emph{Annu. Rev. Fluid Mech.}, vol.~28, pp.~477--539, 1996.",
    'gupta': r"S. Gupta, J. Zhao, A. Sharma, A. Agrawal, K. Hourigan, and M.\,C. Thompson,\\"
             r"``Two- and three-dimensional wake transitions of a NACA0012 airfoil,''\\"
             r"\emph{Journal of Fluid Mechanics}, vol.~954, A26, 2023.",
    'durante': r"D. Durante, E. Rossi, and A. Colagrossi,\\"
               r"``Bifurcations and chaos transition of the flow over an airfoil at low Reynolds number varying the angle of attack,''\\"
               r"\emph{Communications in Nonlinear Science and Numerical Simulation}, art.~105285, 2020.",
    'kurtulus': r"D.\,F. Kurtulu\c{s}, ``On the wake pattern of symmetric airfoils for different incidence angles at $Re=1000$,''\\"
                r"\emph{International Journal of Micro Air Vehicles}, vol.~8, no.~2, pp.~109--139, 2016.",
    'feynman': r"R.\,P. Feynman, R.\,B. Leighton, and M. Sands,\\"
               r"\emph{The Feynman Lectures on Physics}, vol.~II, ch.~41, ``The Flow of Wet Water,''\\"
               r"Addison-Wesley, 1964.",
    'sindy': r"S.\,L. Brunton, J.\,L. Proctor, and J.\,N. Kutz,\\"
             r"``Discovering governing equations from data by sparse identification of nonlinear dynamical systems,''\\"
             r"\emph{Proceedings of the National Academy of Sciences}, vol.~113, no.~15, pp.~3932--3937, 2016.",
    'weller': r"H.\,G. Weller, G. Tabor, H. Jasak, and C. Fureby,\\"
              r"``A tensorial approach to computational continuum mechanics using object-oriented techniques,''\\"
              r"\emph{Computers in Physics}, vol.~12, no.~6, pp.~620--631, 1998.",
}
CITED = []


def cite(key):
    """Number references in order of first appearance: returns (number, png)."""
    if key not in CITED:
        CITED.append(key)
    n = CITED.index(key) + 1
    png = latex_png(f'ref_{key}', r'\color{ink}\begin{tabular}{@{}r@{}}' + f'[{n}]\\ ' + REFS[key]
                    + r'\end{tabular}', math=False)
    return n, png


def reference(slide, path, right=12.73, y=5.8, scale=1.75, max_w=12.1):
    """Place a LaTeX-typeset reference, right-aligned to x = right (inches)."""
    wpx, hpx = Image.open(path).size
    scale = min(scale, max_w * EQ_DPI / wpx)
    w, h = wpx / EQ_DPI * scale, hpx / EQ_DPI * scale
    slide.shapes.add_picture(str(path), Inches(right - w), Inches(y), Inches(w), Inches(h))


# --------------------------------------------------------------- slides
def slide_solver(prs, logo):
    s = new_slide(prs, 'Methodology')
    n, ref_png = cite('weller')
    header(s, 'Numerical Solver')
    table(s, [
        ('Code', f'OpenFOAM, finite volume [{n}]'),
        ('Flow model', '2D incompressible Navier–Stokes, laminar; no turbulence model'),
        ('Pressure–velocity coupling', 'PISO / PIMPLE'),
        ('Convection', 'Second-order linear upwind'),
        ('Time integration', 'Second-order implicit (backward)'),
    ], x=0.6, y=1.75, w=7.6, col0=2.9, size=18, row_h=0.72)
    picture(s, logo, 8.75, 2.35, w=4.0)
    reference(s, ref_png, y=5.8)
    notes(s, pptx_text('solver'))


def slide_mesh(prs, scaling_eq):
    s = new_slide(prs, 'Methodology')
    header(s, 'Computational Mesh')
    picture(s, ROOT / 'figs' / 'mesh_figure.png', 0.6, 1.45, w=8.3)       # h = 3.46
    picture(s, ROOT / 'figs' / 'mesh_bl.png', 9.25, 1.45, w=3.5)          # h = 2.64
    textbox(s, 9.25, 4.12, 3.5, 0.4, 'Boundary-layer resolution, x/c = 0.35–0.65',
            size=11, color=GREY, align=PP_ALIGN.CENTER)
    table(s, [
        ('Total cells', '63,360 (quadrilateral)'),
        ('Domain radius', '30c'),
        ('First-cell height', '2.24 × 10⁻³ c'),
        ('Cells inside δ₉₉ at x = c', '44–49'),
    ], x=0.6, y=5.1, w=6.6, col0=3.3, size=16, row_h=0.42)
    textbox(s, 7.7, 5.05, 5.0, 0.45, 'Blasius flat-plate scaling', size=18, bold=True)
    equation(s, scaling_eq, 7.75, 5.6)
    notes(s, pptx_text('mesh'))


def slide_ale(prs, mp4, poster, eqs):
    s = new_slide(prs, 'Methodology')
    header(s, 'Pitching Motion: ALE Formulation')
    vw = 6.6
    vh = vw * 720 / 1280
    video(s, mp4, poster, 0.6, 1.55, vw, vh)
    textbox(s, 0.6, 1.6 + vh, vw, 0.4, 'The whole O-grid rotates rigidly about the pivot (click to play)',
            size=12, color=GREY, align=PP_ALIGN.CENTER)

    x, y = 7.55, 1.6
    textbox(s, x, y, 4.6, 0.4, 'ALE Navier–Stokes', size=18, bold=True)
    y += 0.5
    y += equation(s, eqs['mom'], x, y) + 0.12
    y += equation(s, eqs['cont'], x, y) + 0.45
    textbox(s, x, y, 4.6, 0.4, 'Mesh velocity (rigid rotation)', size=18, bold=True)
    y += 0.5
    y += equation(s, eqs['um'], x, y) + 0.1
    equation(s, eqs['omega'], x, y)
    notes(s, pptx_text('ale'))



THESIS_TITLE = ('Modeling and Control of Unsteady Flows over Rapidly Maneuvering Airfoils '
                'at Ultra-Low Reynolds Number')


def slide_cover(prs):
    s = new_slide(prs, None)
    picture(s, ROOT / 'buet_logo.jpg', 13.333 / 2 - 0.6, 0.35, h=1.2)
    textbox(s, 0.9, 1.7, 11.5, 0.45, 'THESIS PREDEFENSE', size=16, bold=True, color=ACCENT,
            align=PP_ALIGN.CENTER)
    textbox(s, 0.9, 2.15, 11.5, 1.8, THESIS_TITLE, size=32, bold=True, align=PP_ALIGN.CENTER)
    rule(s, 13.333 / 2 - 0.7, 4.2)
    textbox(s, 2.2, 4.45, 4.3, 0.9, 'Imdadullah Raji\n2110099', size=22, align=PP_ALIGN.CENTER)
    textbox(s, 6.83, 4.45, 4.3, 0.9, 'Abdullah Al Mamun\n2110062', size=22, align=PP_ALIGN.CENTER)
    textbox(s, 0.9, 5.5, 11.5, 0.5, 'Supervisor: Dr Md Ali', size=22, bold=True, align=PP_ALIGN.CENTER)
    textbox(s, 0.9, 6.2, 11.5, 0.8, 'Department of Mechanical Engineering\n'
            'Bangladesh University of Engineering and Technology', size=16, align=PP_ALIGN.CENTER)
    notes(s, pptx_text('cover'))


def slide_intro(prs, mp4, poster, quote, ns):
    s = new_slide(prs, '')
    n, ref_png = cite('feynman')
    header(s, THESIS_TITLE, size=28, rule_y=1.45)
    wpx, hpx = Image.open(quote).size
    qw = wpx / EQ_DPI * 1.75
    picture(s, quote, 0.6, 1.95, w=qw)
    y = 1.95 + hpx / EQ_DPI * 1.75 + 0.35
    y += equation(s, ns['mom'], 0.75, y, scale=2.2) + 0.08
    equation(s, ns['cont'], 0.75, y, scale=2.2)
    vw = 6.2
    vh = vw * 368 / 608
    video(s, mp4, poster, 6.53, 1.7, vw, vh)
    textbox(s, 6.53, 1.72 + vh, vw, 0.4, 'Leading-edge vortex roll-up, Re = 500 (click to play)',
            size=13, color=GREY, align=PP_ALIGN.CENTER)
    reference(s, ref_png, y=5.95)
    notes(s, pptx_text('opening'))


KURTULUS_MODES = ['Attached flow', 'Von Kármán-like street', 'Alternating vortex pairs',
                  'Alternate vortex + single vortex', 'Bluff-body shedding']


def slide_wake_modes(prs, panels, videos):
    s = new_slide(prs, 'Literature Review')
    n, ref_png = cite('kurtulus')
    header(s, 'Wake Patterns at Ultra-Low Reynolds Number')
    textbox(s, 0.6, 1.3, 12, 0.4, f'Kurtuluş [{n}]: five wake modes, Re = 1000', size=16, bold=True, color=ACCENT)
    # equal heights, widths follow each figure's own aspect ratio
    ratios = [Image.open(p).size[0] / Image.open(p).size[1] for p in panels]
    gap, total = 0.12, 12.13
    h = (total - gap * (len(panels) - 1)) / sum(ratios)
    x = 0.6
    for k, (p, r) in enumerate(zip(panels, ratios), start=1):
        picture(s, p, x, 1.72, w=r * h)
        textbox(s, x, 1.74 + h, r * h, 0.55, f'Mode {k}\n{KURTULUS_MODES[k - 1]}', size=13,
                align=PP_ALIGN.CENTER)
        x += r * h + gap
    vw = 5.9
    for i, (mp4, poster, label) in enumerate(videos):
        vx = 0.6 + i * (vw + 0.33)
        vh = vw * Image.open(poster).size[1] / Image.open(poster).size[0]
        textbox(s, vx, 3.4, vw, 0.4, label, size=15, bold=True, color=ACCENT)
        video(s, mp4, poster, vx, 3.8, vw, vh)
    reference(s, ref_png, y=6.2)
    notes(s, pptx_text('wake'))


def slide_durante(prs, panels):
    s = new_slide(prs, 'Literature Review')
    n, ref_png = cite('durante')
    header(s, 'Period Doubling, Tripling and Chaos')
    textbox(s, 0.6, 1.3, 12, 0.4, f'Durante et al. [{n}]: lift at Re = 1000', size=16, bold=True, color=ACCENT)
    cols = [('p1', 'Period 1, α = 13°'), ('p2', 'Period 2, α = 22°'),
            ('p3', 'Period 3, α = 25°'), ('chaos', 'Chaotic, α = 27°')]
    rows = [('history', 'Lift history', 1.95, 0.85), ('spectrum', 'Spectrum', 2.95, 1.25),
            ('portrait', 'Phase portrait', 4.35, 1.3)]
    x0, gap = 1.95, 0.14
    cw = (12.73 - x0 - gap * 3) / 4
    for j, (regime, title) in enumerate(cols):
        textbox(s, x0 + j * (cw + gap), 1.6, cw, 0.35, title, size=15, bold=True, align=PP_ALIGN.CENTER)
    for row, label, y, h in rows:
        textbox(s, 0.6, y + h / 2 - 0.2, 1.3, 0.4, label, size=14, bold=True, color=ACCENT)
        for j, (regime, _) in enumerate(cols):
            png = panels[regime, row]
            wpx, hpx = Image.open(png).size
            w = min(cw, h * wpx / hpx)
            picture(s, png, x0 + j * (cw + gap) + (cw - w) / 2, y, w=w)
    reference(s, ref_png, y=5.85)
    notes(s, pptx_text('durante'))


WILLIAMSON = Path('/home/raji/Research/Notes/airfoil_research/Figures')


def slide_3d(prs, regime_map):
    s = new_slide(prs, 'Literature Review')
    nw, ref_w = cite('williamson')
    ng, ref_g = cite('gupta')
    header(s, 'Beyond 2D: Spanwise Instability')
    textbox(s, 0.6, 1.3, 6.8, 0.4, f'Gupta et al. [{ng}]: NACA0012 wake regimes', size=16, bold=True, color=ACCENT)
    picture(s, regime_map, 0.6, 1.85, w=6.8)
    textbox(s, 7.65, 1.3, 5.1, 0.4, f'Cylinder wake, Williamson [{nw}]', size=16, bold=True, color=ACCENT)
    y = 1.8
    for img, label in [('williamsonModeAB.png', 'Mode A\nRe = 200\nλ ≈ 4D'),
                       ('williamsonModeC.png', 'Mode B\nRe = 270\nλ ≈ 1D')]:
        pic = picture(s, WILLIAMSON / img, 7.65, y, h=1.8)
        textbox(s, 7.65 + pic.width / 914400 + 0.12, y + 0.45, 1.6, 1.0, label, size=15)
        y += 1.9
    reference(s, ref_w, y=5.7)
    reference(s, ref_g, y=6.0)
    notes(s, pptx_text('spanwise'))


def slide_lev(prs, mp4, poster, eqs):
    s = new_slide(prs, 'Literature Review')
    n, ref_png = cite('eldredge2019')
    header(s, 'Pitching Airfoil and the Leading-Edge Vortex')
    vw = 6.9
    vh = vw * 1052 / 2412
    textbox(s, 0.6, 1.3, vw, 0.4, 'Present simulation: Re = 500, K = 0.4, α → 30°', size=16,
            bold=True, color=ACCENT)
    video(s, mp4, poster, 0.6, 1.75, vw, vh)
    textbox(s, 0.6, 1.77 + vh, vw, 0.35, 'Added-mass and circulatory lift (click to play)', size=13, color=GREY, align=PP_ALIGN.CENTER)
    y = 2.18 + vh
    textbox(s, 0.6, y, 2.3, 0.4, 'Kutta–Joukowski', size=16, bold=True, color=ACCENT)
    equation(s, eqs['kj'], 0.75, y + 0.45, scale=2.2)
    textbox(s, 3.0, y, 4.0, 0.4, 'Added mass (Theodorsen)', size=16, bold=True, color=ACCENT)
    equation(s, eqs['am'], 3.15, y + 0.38, scale=2.2)
    textbox(s, 3.0, y + 1.1, 3.75, 0.6, 'a: pivot aft of mid-chord in half-chords, −½ here\n'
            'α̇, α̈ in units of U∞/c, U∞²/c²', size=12, color=GREY)
    textbox(s, 7.85, 1.3, 4.9, 0.4, f'Eldredge & Jones [{n}]: pitching plate', size=16,
            bold=True, color=ACCENT)
    picture(s, ROOT / 'assets' / 'pitching_plate.png', 7.95, 1.75, w=4.78)
    textbox(s, 7.95, 5.47, 4.78, 0.35, 'Faster pitch: tighter vortex, higher lift', size=13, color=GREY,
            align=PP_ALIGN.CENTER)
    reference(s, ref_png, y=6.05)
    notes(s, pptx_text('lev'))


def slide_rom(prs, mp4, poster):
    s = new_slide(prs, 'Objective')
    header(s, 'Objective: Reduced-Order Models')
    vw = 8.3
    vh = vw * 9 / 16
    video(s, mp4, poster, 0.45, 1.45, vw, vh)
    textbox(s, 0.45, 1.5 + vh, vw, 0.4, 'POD of the Re = 500, α = 28° wake (click to play)',
            size=13, color=GREY, align=PP_ALIGN.CENTER)
    bullets(s, 9.05, 1.75, 3.9, 4.8, [
        'Flow control must act within fractions of a millisecond',
        'Goal: a simplified, reduced-order model',
        'POD: the flow as a sum of its most energetic modes',
        '16 GB of simulation data → 12 MB of modes',
        '4-mode reconstruction: ≈ 20 µs',
    ], size=20)
    notes(s, pptx_text('rom'))


def slide_control(prs, eqs):
    s = new_slide(prs, 'Objective')
    n, ref_png = cite('sindy')
    header(s, 'Objective: Model-Based Active Flow Control')
    y = 1.75
    for label, key in [('Reduced dynamics', 'f'), ('Linear model', 'lin'), ('Controller', 'ctrl')]:
        textbox(s, 0.6, y, 5.5, 0.45, label, size=18, bold=True, color=ACCENT)
        y += 0.5
        y += equation(s, eqs[key], 0.75, y, scale=2.8) + 0.4
    bullets(s, 6.9, 1.8, 5.9, 3.9, [
        'x: reduced flow state (e.g. POD coefficients); u: actuation',
        'Model f: how x evolves under actuation',
        'First test: is a linear model enough?',
        f'Nonlinear dynamics → SINDy: a sparse nonlinear model identified from simulation data [{n}]',
        'Goal: a controller that maps x to u',
    ], size=20)
    reference(s, ref_png, y=5.85)
    notes(s, pptx_text('control'))


def slide_resources(prs, books):
    s = new_slide(prs, 'Resources')
    header(s, 'Books Consulted and Code')
    textbox(s, 0.6, 1.55, 6, 0.5, 'Books', size=22, bold=True, color=ACCENT)
    wpx, hpx = Image.open(books).size
    picture(s, books, 0.75, 2.15, w=min(11.8, wpx / EQ_DPI * 1.75))
    textbox(s, 0.6, 4.55, 6, 0.5, 'Code', size=22, bold=True, color=ACCENT)
    tb = s.shapes.add_textbox(Inches(0.75), Inches(5.1), Inches(11.5), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, url in enumerate(['https://github.com/Imdadullah-Raji/flowkit',
                             'https://github.com/Imdadullah-Raji/openfoam_cases']):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = url
        r.hyperlink.address = url
        r.font.name, r.font.size = FONT, Pt(22)
        r.font.color.rgb = INK
    notes(s, pptx_text('resources'))


def slide_qna(prs):
    s = new_slide(prs, '')
    textbox(s, 0.9, 2.55, 11.5, 1.2, 'Questions?', size=60, bold=True, align=PP_ALIGN.CENTER)
    rule(s, 13.333 / 2 - 0.7, 3.95)
    textbox(s, 0.9, 4.25, 11.5, 0.6, 'Thank you', size=24, align=PP_ALIGN.CENTER)
    notes(s, pptx_text('qna'))


def main():
    MEDIA.mkdir(parents=True, exist_ok=True)
    logo = openfoam_logo()
    mesh_mp4 = mesh_motion_video()
    mesh_poster = FRAMES / 'moveMesh.0000.png'
    intro_mp4 = ROOT / 'assets' / 'intro_viz.mp4'
    intro_poster = MEDIA / 'intro_poster.png'
    run(['ffmpeg', '-y', '-i', str(intro_mp4), '-frames:v', '1', str(intro_poster)])
    pod_mp4 = ROOT / 'figs' / 'pod_video' / 'pod_superposition.mp4'
    pod_poster = ROOT / 'figs' / 'pod_video' / 'pod_superposition.png'

    quote = latex_png('quote_feynman', r'\color{ink}\begin{minipage}{3.1in}\raggedright\large\itshape '
                      r'``The test of science is its ability to predict. Had you never visited the earth, '
                      r'could you predict the thunderstorms, the volcanos, the ocean waves, the auroras, '
                      r"and the colorful sunset?''\par\medskip\upshape\raggedleft --- Richard P. Feynman"
                      r'\end{minipage}', math=False)
    books = latex_png('books', r'\color{ink}\begin{minipage}{6.9in}\begin{enumerate}\setlength\itemsep{2pt}'
                      r'\item S.\,L. Brunton and J.\,N. Kutz, \emph{Data-Driven Science and Engineering: '
                      r'Machine Learning, Dynamical Systems, and Control}, 2nd ed. Cambridge University Press, 2022.'
                      r'\item S.\,H. Strogatz, \emph{Nonlinear Dynamics and Chaos: With Applications to Physics, '
                      r'Biology, Chemistry, and Engineering}, 3rd ed. CRC Press, 2024.'
                      r'\item J. Katz and A. Plotkin, \emph{Low-Speed Aerodynamics}, 2nd ed. '
                      r'Cambridge University Press, 2001.'
                      r'\end{enumerate}\end{minipage}', math=False)
    scaling = latex_png('eq_scaling',
                        r'\frac{h_1}{c}=\frac{0.05}{\sqrt{Re}}\approx\frac{1}{100}\,\frac{\delta_{99}(c)}{c}')
    eqs = {
        'mom': latex_png('eq_mom',
                         r'\left.\frac{\partial\mathbf{u}}{\partial t}\right|_{\boldsymbol\xi}'
                         r'+\big[(\mathbf{u}-\textcolor{accent}{\mathbf{u}_m})\cdot\nabla\big]\mathbf{u}'
                         r'=-\nabla p_k+\nu\nabla^2\mathbf{u}'),
        'cont': latex_png('eq_cont', r'\nabla\cdot\mathbf{u}=0'),
        'um': latex_png('eq_um', r'\textcolor{accent}{\mathbf{u}_m}=\boldsymbol\Omega\times\mathbf{r}'),
        'omega': latex_png('eq_omega', r'\boldsymbol\Omega(t)=\dot\alpha(t)\,\hat{\mathbf{k}}'),
    }
    ctrl = {
        'f': latex_png('eq_f', r'\dot{\mathbf{x}}=\mathbf{f}(\mathbf{x},\mathbf{u},t)'),
        'lin': latex_png('eq_lin', r'\dot{\mathbf{x}}=\mathbf{A}\mathbf{x}+\mathbf{B}\mathbf{u}'),
        'ctrl': latex_png('eq_ctrl', r'\mathbf{u}=\mathcal{K}(\mathbf{x})'),
    }

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_cover(prs)
    ns = {
        'mom': latex_png('eq_ns', r'\frac{\partial\mathbf{u}}{\partial t}+(\mathbf{u}\cdot\nabla)\mathbf{u}'
                         r'=-\frac{1}{\rho}\nabla p+\nu\nabla^2\mathbf{u}'),
        'cont': latex_png('eq_ns_cont', r'\nabla\cdot\mathbf{u}=0'),
    }
    slide_intro(prs, intro_mp4, intro_poster, quote, ns)
    slide_rom(prs, pod_mp4, pod_poster)
    slide_control(prs, ctrl)
    slide_solver(prs, logo)
    slide_mesh(prs, scaling)
    slide_ale(prs, mesh_mp4, mesh_poster, eqs)
    wake = ROOT / 'figs' / 'wake_video'
    slide_wake_modes(prs, kurtulus_panels(), [
        (wake / 're500_aoa14.mp4', wake / 're500_aoa14.png', 'Present simulation, Re = 500, α = 14°: mode 2'),
        (wake / 're500_aoa26.mp4', wake / 're500_aoa26.png', 'Present simulation, Re = 500, α = 26°: mode 3'),
    ])
    slide_durante(prs, durante_panels())
    slide_3d(prs, gupta_map())
    lev_mp4 = ROOT / 'assets' / 'lev_exp_vid.mp4'
    lev_poster = MEDIA / 'lev_poster.png'
    run(['ffmpeg', '-y', '-i', str(lev_mp4), '-frames:v', '1', str(lev_poster)])
    slide_lev(prs, lev_mp4, lev_poster, {
        'kj': latex_png('eq_kj', r"L'=\rho\,U_\infty\,\Gamma"),
        'am': latex_png('eq_added_mass', r'C_{L,\mathrm{nc}}=\frac{\pi}{2}\Big(\dot\alpha-\frac{a}{2}\,\ddot\alpha\Big)'),
    })
    slide_resources(prs, books)
    slide_qna(prs)
    footers(prs)
    out = OUT / 'predefense.pptx'
    prs.save(out)
    print(out)


if __name__ == '__main__':
    main()
